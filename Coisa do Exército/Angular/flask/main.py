from flask import Flask, request, send_file, jsonify
from werkzeug.utils import secure_filename
from docx import Document
import pymupdf
import pandas as pd
import os
from flask_cors import CORS


app = Flask(__name__)
upload_folder = 'uploads'
app.config['UPLOAD_FOLDER'] = upload_folder

CORS(app)

if not os.path.exists(upload_folder):
    os.makedirs(upload_folder)

@app.route('/pesquisa_preco', methods=['POST'])
def uploadPdf():
    if 'pdf_file' not in request.files:
        return "Nenhum arquivo enviado"
    pdf_file = request.files['pdf_file']
    if pdf_file.filename == '':
        return "Arquivo não selecionado"
    if pdf_file:
        file_name = secure_filename(pdf_file.filename)
        pdf_path = os.path.join(app.config['UPLOAD_FOLDER'], file_name)
        pdf_file.save(pdf_path)
        csv_path = extrairMediana(pdf_path)
        return send_file(csv_path, as_attachment=True, download_name='resultado.csv')
    return "Erro ao enviar o arquivo"

@app.post("/gerar_doc")
def gerar_doc():
    dados = request.json  # lista de itens

    doc = Document()
    doc.add_heading("Tabela Gerada", level=1)

    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = 'Descrição'
    hdr[1].text = 'Catmat'
    hdr[2].text = 'Quantidade'
    hdr[3].text = 'Unidade'

    for item in dados:
        row = table.add_row().cells
        row[0].text = str(item.get("descricao", ""))
        row[1].text = str(item.get("catmat", ""))
        row[2].text = str(item.get("quantidade_total", ""))
        row[3].text = str(item.get("unidadeMedida", ""))

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="documento.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )



def extrairMediana(pdf_path):
    dataframe = []
    doc = pymupdf.open(pdf_path)
    page = doc[0]
    for page in doc:
        texto = page.search_for("Consolidação dos preços cotados")
        if texto:
            cropstart = texto
            rx0 = cropstart[0].x0 + 200
            ry0 = cropstart[0].y0 + 40

            rx1 = rx0 + 100
            ry1 = ry0 + 30
            rect = pymupdf.Rect(rx0, ry0, rx1, ry1)
            print(rect)
            mediana_text = page.get_text(clip=rect)
            print("Mediana extraída:", mediana_text)
            mediana_value = mediana_text.strip().replace(',', '.')
            dataframe.append({'Mediana': mediana_value})
            df = pd.DataFrame(dataframe)
        else:
            print("Texto 'Consolidação dos preços cotados' não encontrado na página.")
    output_csv = 'teste.csv'
    df.to_csv(output_csv, index=False)
    return output_csv

if __name__ == "__main__":
    app.run(debug=True)